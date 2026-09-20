from __future__ import annotations

import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.source_registry import SourceRegistryService
from app.store import KnowledgeStore


def _import_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "来源台账"
    sheet.append([
        "来源编码", "国家/地区", "来源名称", "来源类别", "主管机构", "网址",
        "接入方式", "优先级", "状态", "共享范围", "语言", "文件类型",
        "更新频率", "采集范围", "备注", "项目维护说明",
    ])
    sheet.append([
        "DE-MANAGED-TEST", "DE", "德国测试来源更新", "认证", "Test Authority",
        "https://example.com/de-updated", "网页", "A", "已确认", "否", "de,en",
        "html,pdf", "每日", "updated scope", "updated note", "updated owner",
    ])
    sheet.append([
        "FR-MANAGED-TEST", "FR", "法国测试来源", "法规", "French Test Authority",
        "https://example.com/fr", "混合", "B", "调研中", "否", "fr",
        "html,pdf", "每周", "test scope", "", "",
    ])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        store = KnowledgeStore(Path(tmp) / "source_management.db")
        service = SourceRegistryService(store)

        before = service.list_sources()["summary"]
        base_count = before["registered_sources"]

        created = service.create({
            "source_key": "DE-MANAGED-TEST",
            "region_code": "DE",
            "source_name": "德国测试来源",
            "source_type": "certification",
            "authority": "Test Authority",
            "base_url": "https://example.com/de",
            "access_method": "html",
            "priority": "B",
            "status": "research",
            "refresh_policy": "weekly",
            "languages": ["de", "en"],
            "file_types": ["html"],
            "crawl_scope": "test scope",
        })
        assert created["source_key"] == "DE-MANAGED-TEST"
        assert created["languages"] == ["de", "en"]
        assert service.list_sources()["summary"]["registered_sources"] == base_count + 1

        updated = service.update(created["id"], {
            "priority": "A",
            "status": "validated",
            "languages": "de、en、fr",
            "file_types": "html,pdf",
        })
        assert updated["priority"] == "A"
        assert updated["status"] == "validated"
        assert updated["languages"] == ["de", "en", "fr"]
        assert updated["file_types"] == ["html", "pdf"]

        batch = service.batch_update([created["id"]], {"status": "connected", "priority": "C"})
        assert batch["updated"] == 1
        after_batch = service.detail(created["id"])
        assert after_batch["status"] == "connected"
        assert after_batch["priority"] == "C"

        exported = service.export_xlsx()
        assert len(exported) > 1000
        workbook = load_workbook(BytesIO(exported), read_only=True, data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows())]
        assert headers[:4] == ["来源编码", "国家/地区", "来源名称", "来源类别"]
        rows = list(sheet.iter_rows(values_only=True))
        assert any(row[0] == "DE-MANAGED-TEST" for row in rows)

        template = service.export_xlsx(template_only=True)
        template_book = load_workbook(BytesIO(template), read_only=True, data_only=True)
        assert template_book.active.max_row == 1

        imported = service.import_xlsx(_import_workbook())
        assert imported["created"] == 1
        assert imported["updated"] == 1
        assert imported["failed"] == 0
        de = service.by_key("DE-MANAGED-TEST")
        fr = service.by_key("FR-MANAGED-TEST")
        assert de["source_name"] == "德国测试来源更新"
        assert de["status"] == "validated"
        assert fr["region_code"] == "FR"
        assert fr["access_method"] == "mixed"

        active_before_archive = service.list_sources()["summary"]["registered_sources"]
        archived = service.archive(fr["id"], note="测试归档")
        assert archived["status"] == "archived"
        summary = service.list_sources()["summary"]
        assert summary["registered_sources"] == active_before_archive - 1
        assert summary["archived_sources"] >= 1
        assert "测试归档" in archived["owner_note"]

    page = Path("static/sources.html").read_text(encoding="utf-8")
    runtime = Path("app/platform_server.py").read_text(encoding="utf-8")
    workflow = Path(".github/workflows/quality.yml").read_text(encoding="utf-8")

    for label in ("新增来源", "Excel导入", "Excel导出", "下载模板", "批量状态", "批量优先级", "关联采集配置"):
        assert label in page
    for endpoint in (
        "/api/sources/create",
        "/api/sources/update",
        "/api/sources/archive",
        "/api/sources/batch",
        "/api/sources/import",
        "/api/sources/export",
    ):
        assert endpoint in page or endpoint in runtime
        assert endpoint in runtime
    assert "collection_profiles" in runtime
    assert "Stage30 source website management tests" in workflow

    print("OK: stage30 source website management CRUD, Excel and batch operations passed")


if __name__ == "__main__":
    main()
