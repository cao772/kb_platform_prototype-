from __future__ import annotations

import json
import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.regions import TARGET_MARKETS
from app.source_registry import SOURCE_CATALOG_PATH, SourceRegistryService
from app.store import KnowledgeStore


def main() -> None:
    catalog = json.loads(SOURCE_CATALOG_PATH.read_text(encoding="utf-8"))
    assert isinstance(catalog, list)
    assert len(catalog) == 258
    assert len({item["source_key"] for item in catalog}) == 258
    assert all(str(item.get("base_url") or "").startswith(("http://", "https://")) for item in catalog)
    assert all(item.get("status") == "research" for item in catalog)

    with tempfile.TemporaryDirectory() as tmp:
        store = KnowledgeStore(Path(tmp) / "source_catalog_300.db")
        service = SourceRegistryService(store)
        listing = service.list_sources(limit=3000)
        items = listing["items"]
        summary = listing["summary"]

        assert summary["target_sources"] == 300
        assert summary["registered_sources"] == 300
        assert summary["remaining_to_target"] == 0
        assert summary["target_markets"] == 21
        assert summary["target_markets_with_sources"] == 21
        assert len(items) == 300
        assert len({item["source_key"] for item in items}) == 300

        target_codes = {market.code for market in TARGET_MARKETS}
        per_region = {}
        for item in items:
            per_region[item["region_code"]] = per_region.get(item["region_code"], 0) + 1
            assert item["source_summary"].strip(), item["source_key"]
            assert item["extractable_summary"].strip(), item["source_key"]
            assert item["base_url"].startswith(("http://", "https://")), item["source_key"]
        assert per_region["EU"] == 20
        assert target_codes <= set(per_region)
        assert min(per_region[code] for code in target_codes) >= 13

        assert summary["by_type"]["regulation"] >= 40
        assert summary["by_type"]["standard"] >= 25
        assert summary["by_type"]["certification"] >= 25
        assert summary["by_type"]["gma"] >= 100

        germany = service.list_sources(region_code="DE", limit=100)
        assert len(germany["items"]) == 14
        assert any(item["source_key"] == "DE-BNETZA-MARKET" for item in germany["items"])
        assert any("市场监管" in item["extractable_summary"] or "无线电" in item["extractable_summary"] for item in germany["items"])

        exported = service.export_xlsx()
        workbook = load_workbook(BytesIO(exported), read_only=True, data_only=True)
        headers = [cell.value for cell in next(workbook.active.iter_rows())]
        assert "网站简述" in headers
        assert "可提取资料简述" in headers
        assert workbook.active.max_row == 301

    page = Path("static/sources.html").read_text(encoding="utf-8")
    runtime = Path("app/source_registry.py").read_text(encoding="utf-8")
    workflow = Path(".github/workflows/quality.yml").read_text(encoding="utf-8")
    assert "网站简述" in page
    assert "可提取资料简述" in page
    assert "source_catalog_258.json" in runtime
    assert "Stage31 300-source catalog tests" in workflow

    print("OK: stage31 300-source catalog and per-source descriptions passed")


if __name__ == "__main__":
    main()
