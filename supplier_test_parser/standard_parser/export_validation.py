"""Validate persisted customer cells, not only normalized OCR text."""
import re

BAD_PATTERNS=(r'士\s*[0-9]',r'1。。',r'脱扣电流1，',r'1、和1，',r'和（依次',r'45\s*通过视检',r'85\s*注1',r'IEC\s*60695-212',r'675\s*一按照',r'2024-0814-1127-0806-4238',r'驻立式1类器具')

def bad_cells(sheets):
    result=[]
    for sheet in sheets:
        for ri,row in enumerate(sheet['rows'],1):
            for ci,value in enumerate(row,1):
                for pattern in BAD_PATTERNS:
                    if re.search(pattern,str(value)):
                        result.append(dict(sheet=sheet['name'],row=ri,column=ci,pattern=pattern))
    return result

def verify_export(path,spec):
    import openpyxl
    wb=openpyxl.load_workbook(path,data_only=False)
    assert wb.sheetnames==[s['name'] for s in spec['sheets']]
    actual=[]
    for s in spec['sheets']:
        w=wb[s['name']];rows=list(w.values)
        assert len(rows)==len(s['rows'])
        for ri,row in enumerate(s['rows'],1):
            for ci,v in enumerate(row,1):
                cell=w.cell(ri,ci)
                assert cell.data_type!='e',(w.title,cell.coordinate)
                if ri>1 and s['rows'][0][ci-1]=='原页链接' and v:
                    assert cell.value==f'=HYPERLINK("{v}","查看原页")'
                else:assert ('' if cell.value is None else cell.value)==('' if v is None else v),(w.title,cell.coordinate)
        actual.append(dict(name=w.title,rows=rows))
    hits=bad_cells(actual)
    assert not hits,hits[:20]
    return {s['name']:len(s['rows'])-1 for s in spec['sheets']}
