"""Validate the exported workbooks and assemble a bounded, secret-free handoff."""
import argparse,hashlib,json,zipfile
from pathlib import Path

def main():
    import openpyxl
    parser=argparse.ArgumentParser();parser.add_argument('output');parser.add_argument('--source-commit',required=True)
    args=parser.parse_args();root=Path(args.output).resolve()
    specs=json.loads((root/'excel_generator/workbooks.json').read_text(encoding='utf8'))
    checks=[]
    html=(root/'原文对照.html').read_text(encoding='utf8')
    for spec in specs:
        workbook=openpyxl.load_workbook(root/spec['filename'],data_only=False)
        assert workbook.sheetnames==[s['name'] for s in spec['sheets']]
        for s in spec['sheets']:
            w=workbook[s['name']]
            assert (w.max_row,w.max_column)==(len(s['rows']),len(s['rows'][0]))
            assert w.freeze_panes=='B2'
            assert len(w.tables)==1
            for row in w:
                for cell in row:
                    if cell.data_type=='e':raise ValueError(f'Excel error {w.title}!{cell.coordinate}')
                    if cell.data_type=='f':
                        assert str(cell.value).startswith('=HYPERLINK(')
                        url=str(cell.value).split('"')[1];anchor=url.split('#')[1]
                        assert f'id="{anchor}"' in html,anchor
            checks.append({'workbook':spec['filename'],'sheet':w.title,'rows':w.max_row-1,'columns':w.max_column,'frozen_headers':True,'filter':True})
    manifest=json.loads((root/'manifest.json').read_text(encoding='utf8'))
    assert sum(d['pdf_pages'] for d in manifest['documents'])==sum(d['ocr_pages'] for d in manifest['documents'])
    assert manifest['model_api_calls']==0
    for d in manifest['documents']:
        nodes=json.loads((root/'clause_tree'/f"{d['code']}.json").read_text(encoding='utf8'))['nodes']
        numbers={n['number'] for n in nodes if n['kind']=='chapter' and not n['chapter'].startswith('附录')}
        assert numbers=={str(i) for i in range(1,33)},(d['code'],numbers)
        for page in range(1,d['pdf_pages']+1):
            assert (root/'ocr'/d['code']/f'{page:03d}.json').exists()
            assert (root/'原页'/f"{d['code']}-{page:03d}.jpg").exists()
    (root/'validation/export_checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf8')
    note='''# 标准解析工程交接

本包仅包含两份GB/T标准，不包含认证规则正式表，也不覆盖业务方的证据抽检工作簿。

打开01_标准全文解析结果.xlsx查看客户31字段及8个主题表；standard_validation_report.xlsx为内部检查报告。完整解压后，Excel中的“查看原页”链接与原文对照.html配套使用。

148页与52页全部处理，两份标准正文第1至32章均有节点。章节目录同时保留原目录摘录和正文结构。正文结构包含章、条、附录，节点数不等于准确条款数；图表索引区分标题与引用。

长单元格分为多行，“内容分段”标示顺序，同编号按顺序拼接可恢复完整字段。extract/standard_records.json保留不拆行的原始记录。ocr/保留识别块与位置；坐标是左下原点归一化xywh。页码一律为PDF物理页码。

本次全部使用本地Apple Vision和原句规则摘录，没有新增模型API调用。样品、周期、条件等字段是原句候选，不代替业务语义签审。第2章标准名称不作为测试要求。未规定字段不按常识补齐；通用条件引用有明确说明，不自动覆盖特殊要求。

内部报告中的编号跳号候选、低置信文字、重复标题候选需要业务方按原页核验。特殊要求101起编号不自动认定为缺条款。8个文字稀疏页已看原图确认为仅水印的空白页。尚未计算准确率，不把处理覆盖率当准确率。

已修订的原图抽查字符保留在validation/applied_corrections.json；未逐条完成业务签审。最终客户版由业务方统一口径、抽检和收口。
'''
    (root/'交接说明.md').write_text(note,encoding='utf8')
    paths=[root/s['filename'] for s in specs]+[root/n for n in ['原文对照.html','manifest.json','交接说明.md','validation/issues.json','validation/applied_corrections.json','validation/export_checks.json']]
    if (root/'checked_corrections.json').exists():paths.append(root/'checked_corrections.json')
    for folder in ['ocr','clause_tree','extract','原页']:paths.extend(p for p in (root/folder).rglob('*') if p.is_file())
    paths=sorted(set(paths));files=[{'path':str(p.relative_to(root)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'size_bytes':p.stat().st_size} for p in paths]
    handoff={'logical_key':'supplier-part2-local-standard-parser','local_path':str(root),'folder_key':'evaluation_data','source_commit':args.source_commit,'artifact_type':'evidence','notes':'仅标准工程解析；不包含02认证规则及03业务抽检表；业务签审另行进行','files':files}
    (root/'handoff_manifest.json').write_text(json.dumps(handoff,ensure_ascii=False,indent=2),encoding='utf8')
    archive=root/'标准解析_工程交接包.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for p in paths+[root/'handoff_manifest.json']:z.write(p,p.relative_to(root))
    with zipfile.ZipFile(archive) as z:assert z.testzip() is None
    print(json.dumps({'validated_sheets':len(checks),'files':len(paths)+1,'archive':str(archive),'bytes':archive.stat().st_size,'sha256':hashlib.sha256(archive.read_bytes()).hexdigest()},ensure_ascii=False))

if __name__=='__main__':main()
