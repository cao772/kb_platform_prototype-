from __future__ import annotations

import io
import json
import os
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from app.docx_parser import parse_docx
from app.model_gateway import call_vision_ocr, current_model_config
from app.pdf_structure import extract_pdf_structure
from app.local_ocr import configuration as local_ocr_configuration, recognize as local_ocr_recognize
from app.runtime_settings import current_parser_settings
from app.vendor_path import activate_vendor

activate_vendor()


@dataclass
class StandardBlock:
    block_type: str
    text: str
    page_no: int | None = None
    section_title: str | None = None
    metadata: dict = field(default_factory=dict)


@dataclass
class StandardDocument:
    source_path: str
    filename: str
    title: str
    mime_type: str
    parser: str
    blocks: list[StandardBlock]
    metadata: dict = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        return "\n".join(block.text for block in self.blocks if block.text.strip())


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.main_parts: list[str] = []
        self.title_parts: list[str] = []
        self._stack: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag not in {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}:
            self._stack.append(tag)

    def handle_endtag(self, tag):
        if tag in self._stack:
            index = len(self._stack) - 1 - self._stack[::-1].index(tag)
            del self._stack[index:]

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if "title" in self._stack and text:
            self.title_parts.append(text)
        if any(tag in self._stack for tag in ("script", "style", "noscript", "svg", "head", "nav", "footer", "header")):
            return
        if text:
            self.parts.append(text)
            if "main" in self._stack:
                self.main_parts.append(text)


def _parser_backend() -> str:
    return (os.getenv("KB_PARSER_BACKEND", "") or current_parser_settings().get("backend") or "lightweight").lower()


def _ocr_enabled() -> bool:
    mode = str(current_parser_settings().get("ocr_mode") or "auto").lower()
    if mode == "off":
        return False
    return (mode == "auto" and bool(local_ocr_configuration())) or current_model_config("vision").configured


def recognize_image(image_bytes: bytes, mime_type: str, *, instruction: str | None = None):
    if current_parser_settings().get("ocr_mode") == "auto" and local_ocr_configuration():
        return local_ocr_recognize(image_bytes)
    return call_vision_ocr(image_bytes, mime_type, instruction=instruction)


def parse_file(path: str | Path) -> StandardDocument:
    file_path = Path(path)
    if _parser_backend() == "docling":
        try:
            return parse_docling_standard(file_path)
        except Exception:
            pass
    suffix = file_path.suffix.lower()
    if suffix == ".docx":
        return parse_docx_standard(file_path)
    if suffix == ".pdf":
        return parse_pdf_standard(file_path)
    if suffix == ".xlsx":
        return parse_xlsx_standard(file_path)
    if suffix == ".pptx":
        return parse_pptx_standard(file_path)
    if suffix in {".txt", ".md", ".csv"}:
        return parse_text_standard(file_path)
    if suffix == ".json":
        return parse_json_standard(file_path)
    if suffix == ".xml":
        return parse_xml_standard(file_path)
    if suffix in {".html", ".htm"}:
        return parse_html_standard(file_path)
    if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
        return parse_image_standard(file_path)
    raise ValueError(f"Unsupported file type: {file_path.suffix}")


def parse_docx_standard(path: Path) -> StandardDocument:
    parsed = parse_docx(path)
    blocks = [StandardBlock("paragraph", text) for text in parsed.paragraphs]
    for index, table in enumerate(parsed.tables, 1):
        blocks.append(StandardBlock(
            "table",
            "\n".join(" | ".join(cell for cell in row if cell) for row in table),
            metadata={"table_index": index},
        ))
    return StandardDocument(
        str(path.resolve()), path.name, parsed.title,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx-xml", blocks,
    )


def _render_pdf_page(path: Path, page_index: int) -> bytes | None:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return None
    doc = fitz.open(str(path))
    try:
        page = doc.load_page(page_index)
        pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
        return pix.tobytes("png")
    finally:
        doc.close()


def parse_pdf_standard(path: Path) -> StandardDocument:
    max_vision_pages = int(current_parser_settings().get("vision_max_pages") or 6)
    ocr_available = _ocr_enabled()

    def _ocr_callback(image_bytes: bytes, page_no: int) -> tuple[str, dict]:
        return recognize_image(
            image_bytes,
            "image/png",
            instruction=(
                f"这是PDF第{page_no}页。请识别正文、章节/条款编号、表格文字、日期、法规标准编号和认证要求。"
                "保持阅读顺序和层级；表格尽量按行列输出；无法确认的字符不要猜测。"
            ),
        )

    structured_blocks, parse_report = extract_pdf_structure(
        path,
        ocr_enabled=ocr_available,
        max_vision_pages=max_vision_pages,
        ocr_callback=_ocr_callback if ocr_available else None,
    )
    blocks = [
        StandardBlock(
            str(item.get("type") or "paragraph"),
            str(item.get("text") or ""),
            page_no=int(item.get("page_no") or 0) or None,
            section_title=str(item.get("section_title") or "") or None,
            metadata={
                "source": "structured_pdf",
                "bbox": item.get("bbox", []),
                "heading_level": item.get("heading_level"),
                "hierarchy": item.get("hierarchy", []),
                "rows": item.get("rows", []) if item.get("type") == "table" else [],
            },
        )
        for item in structured_blocks
        if str(item.get("text") or "").strip()
    ]
    summary = dict(parse_report.get("summary") or {})
    if not blocks:
        blocks = [StandardBlock("page_unreadable", "[PDF未识别到可索引内容]", metadata={"source": "unreadable"})]
    parser_name = "pymupdf-layout+vision" if summary.get("ocr_pages") else "pymupdf-layout"
    if summary.get("ocr_pages") and any("local_ocr" in page.get("strategy", "") for page in parse_report.get("pages", [])):
        parser_name = "pymupdf-layout+paddleocr"
    return StandardDocument(
        str(path.resolve()), path.name, path.stem, "application/pdf", parser_name,
        blocks,
        {
            "page_count": int(summary.get("page_count") or 0),
            "empty_text_pages": int(summary.get("scan_pages") or 0),
            "vision_ocr_pages": int(summary.get("ocr_pages") or 0),
            "ocr_available": ocr_available,
            "parse_summary": summary,
            "parse_report": parse_report,
        },
    )


def parse_xlsx_standard(path: Path) -> StandardDocument:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    blocks: list[StandardBlock] = []
    total_rows = 0
    for sheet in wb.worksheets:
        lines = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_index > 2000:
                lines.append("[工作表内容过长，已截断]")
                break
            cells = [str(value).strip() for value in row[:100] if value is not None and str(value).strip()]
            if cells:
                lines.append(" | ".join(cells))
                total_rows += 1
        if lines:
            blocks.append(StandardBlock(
                "sheet",
                "\n".join(lines),
                section_title=sheet.title,
                metadata={"sheet_name": sheet.title, "non_empty_rows": len(lines)},
            ))
    return StandardDocument(
        str(path.resolve()), path.name, path.stem,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "openpyxl", blocks, {"sheet_count": len(wb.sheetnames), "non_empty_rows": total_rows},
    )


def parse_pptx_standard(path: Path) -> StandardDocument:
    from pptx import Presentation
    deck = Presentation(str(path))
    blocks: list[StandardBlock] = []
    for index, slide in enumerate(deck.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                texts.append(text.strip())
        if texts:
            blocks.append(StandardBlock(
                "slide",
                "\n".join(texts),
                page_no=index,
                metadata={"slide_index": index},
            ))
    return StandardDocument(
        str(path.resolve()), path.name, path.stem,
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "python-pptx", blocks, {"slide_count": len(deck.slides)},
    )


def parse_text_standard(path: Path) -> StandardDocument:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return StandardDocument(
        str(path.resolve()), path.name, path.stem, "text/plain", "plain-text",
        [StandardBlock("text", text)],
    )


def parse_json_standard(path: Path) -> StandardDocument:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    try:
        text = json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        text = raw
    return StandardDocument(
        str(path.resolve()), path.name, path.stem, "application/json", "json",
        [StandardBlock("json", text)],
    )


def parse_xml_standard(path: Path) -> StandardDocument:
    import xml.etree.ElementTree as ET

    raw = path.read_text(encoding="utf-8", errors="replace")
    try:
        root = ET.fromstring(path.read_bytes())
    except ET.ParseError:
        return StandardDocument(
            str(path.resolve()), path.name, path.stem, "application/xml", "xml-raw",
            [StandardBlock("xml_text", raw)],
            {"well_formed": False},
        )

    lines: list[str] = []
    for fragment in root.itertext():
        text = " ".join(fragment.split())
        if not text:
            continue
        lines.append(text)
    return StandardDocument(
        str(path.resolve()), path.name, path.stem, "application/xml", "xml-etree",
        [StandardBlock("xml_text", "\n".join(lines))],
        {"root_tag": str(root.tag).split("}")[-1], "well_formed": True, "element_count": sum(1 for _ in root.iter())},
    )


def parse_html_standard(path: Path) -> StandardDocument:
    parser = _HTMLTextExtractor()
    raw = path.read_bytes()
    charset = re.search(br'charset\s*=\s*["\x27]?([A-Za-z0-9_-]+)', raw[:8192], re.I)
    encoding = charset.group(1).decode('ascii') if charset else 'utf-8'
    try:
        text = raw.decode(encoding, errors='replace')
    except LookupError:
        text = raw.decode('utf-8', errors='replace')
    parser.feed(text)
    return StandardDocument(
        str(path.resolve()), path.name, " ".join(parser.title_parts) or path.stem, "text/html", "html-parser",
        [StandardBlock("html_text", "\n".join(parser.main_parts or parser.parts))],
        {"encoding": encoding, "main_content_selected": bool(parser.main_parts)},
    )


def parse_image_standard(path: Path) -> StandardDocument:
    from PIL import Image

    image = Image.open(path)
    mime_type = f"image/{'jpeg' if path.suffix.lower() in {'.jpg', '.jpeg'} else path.suffix.lower().lstrip('.')}"
    metadata = {
        "width": image.width,
        "height": image.height,
        "mode": image.mode,
    }
    if _ocr_enabled():
        raw = path.read_bytes()
        recognized, trace = recognize_image(raw, mime_type)
        if recognized:
            metadata.update({"ocr_status": "completed", "ocr_trace": trace})
            return StandardDocument(
                str(path.resolve()), path.name, path.stem, mime_type, "vision-model",
                [StandardBlock("image_ocr", recognized, metadata=metadata)], metadata,
            )
        metadata.update({"ocr_status": "failed", "ocr_trace": trace})
    else:
        metadata["ocr_status"] = "model_not_configured"
    text = f"[图片待识别] {path.name} {image.width}x{image.height}"
    return StandardDocument(
        str(path.resolve()), path.name, path.stem, mime_type, "image-metadata",
        [StandardBlock("image", text, metadata=metadata)], metadata,
    )


def parse_docling_standard(path: Path) -> StandardDocument:
    from docling.document_converter import DocumentConverter
    result = DocumentConverter().convert(str(path))
    markdown = result.document.export_to_markdown()
    return StandardDocument(
        str(path.resolve()), path.name, path.stem, "application/octet-stream", "docling",
        [StandardBlock("docling_document", markdown, metadata={"structure_aware": True})],
        {"docling": True},
    )
