from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from collections import Counter
from dataclasses import dataclass, asdict
from typing import Any

from openpyxl import Workbook, load_workbook

from app.regions import TARGET_MARKETS, canonical_region_code
from app.store import KnowledgeStore

TARGET_SOURCE_COUNT = 300
SOURCE_TYPES = {"regulation", "standard", "certification", "gma"}
SOURCE_STATUSES = {"research", "validated", "connected", "paused", "archived"}
SOURCE_PRIORITIES = {"A", "B", "C"}
ACCESS_METHODS = {"html", "pdf", "xml", "api", "rss", "mixed", "metadata_only"}
REFRESH_POLICIES = {"daily", "weekly", "monthly", "event", "manual"}
HARVESTABILITY_TYPES = {"direct", "adapter", "metadata_only", "manual_authorized", "unavailable"}
VERIFICATION_STATUSES = {"pending", "preclassified", "verified", "failed"}

SOURCE_EXPORT_COLUMNS = (
    ("来源编码", "source_key"),
    ("国家/地区", "region_code"),
    ("来源名称", "source_name"),
    ("来源类别", "source_type"),
    ("主管机构", "authority"),
    ("网址", "base_url"),
    ("接入方式", "access_method"),
    ("优先级", "priority"),
    ("状态", "status"),
    ("共享范围", "shared_scope"),
    ("语言", "languages"),
    ("文件类型", "file_types"),
    ("更新频率", "refresh_policy"),
    ("可采集性", "harvestability"),
    ("核验状态", "verification_status"),
    ("最近核验时间", "last_verified_at"),
    ("HTTP状态", "last_http_status"),
    ("内容类型", "last_content_type"),
    ("核验说明", "verification_note"),
    ("网站简述", "source_summary"),
    ("可提取资料简述", "extractable_summary"),
    ("采集范围", "crawl_scope"),
    ("备注", "notes"),
    ("项目维护说明", "owner_note"),
)

SOURCE_HEADER_ALIASES = {
    "来源编码": "source_key", "source_key": "source_key",
    "国家/地区": "region_code", "region_code": "region_code",
    "来源名称": "source_name", "source_name": "source_name",
    "来源类别": "source_type", "source_type": "source_type",
    "主管机构": "authority", "authority": "authority",
    "网址": "base_url", "base_url": "base_url", "url": "base_url",
    "接入方式": "access_method", "access_method": "access_method",
    "优先级": "priority", "priority": "priority",
    "状态": "status", "status": "status",
    "共享范围": "shared_scope", "shared_scope": "shared_scope",
    "语言": "languages", "languages": "languages",
    "文件类型": "file_types", "file_types": "file_types",
    "更新频率": "refresh_policy", "refresh_policy": "refresh_policy",
    "可采集性": "harvestability", "harvestability": "harvestability",
    "核验状态": "verification_status", "verification_status": "verification_status",
    "最近核验时间": "last_verified_at", "last_verified_at": "last_verified_at",
    "HTTP状态": "last_http_status", "last_http_status": "last_http_status",
    "内容类型": "last_content_type", "last_content_type": "last_content_type",
    "核验说明": "verification_note", "verification_note": "verification_note",
    "网站简述": "source_summary", "source_summary": "source_summary",
    "可提取资料简述": "extractable_summary", "extractable_summary": "extractable_summary",
    "采集范围": "crawl_scope", "crawl_scope": "crawl_scope",
    "备注": "notes", "notes": "notes",
    "项目维护说明": "owner_note", "owner_note": "owner_note",
}

SOURCE_TYPE_ALIASES = {"法规": "regulation", "标准": "standard", "认证": "certification", "GMA/市场准入": "gma", "GMA": "gma"}
SOURCE_STATUS_ALIASES = {"调研中": "research", "已确认": "validated", "已接入": "connected", "暂停": "paused", "已归档": "archived"}
ACCESS_METHOD_ALIASES = {"网页": "html", "PDF": "pdf", "XML": "xml", "API": "api", "RSS": "rss", "混合": "mixed", "仅元数据": "metadata_only"}
REFRESH_POLICY_ALIASES = {"每日": "daily", "每周": "weekly", "每月": "monthly", "按事件": "event", "人工": "manual"}
SOURCE_CATALOG_PATH = Path(__file__).resolve().parent.parent / "data" / "source_catalog_258.json"


def _default_source_summary(*, name: str, authority: str, source_type: str, crawl_scope: str) -> str:
    type_label = {"regulation": "法规", "standard": "标准", "certification": "认证", "gma": "市场准入"}.get(source_type, "法规认证")
    authority_text = authority or "相关主管机构"
    scope = str(crawl_scope or "").strip()
    tail = f"，重点覆盖{scope}" if scope else ""
    return f"{name}是{authority_text}维护的{type_label}相关官方来源{tail}。"


def _default_extractable_summary(*, source_type: str, crawl_scope: str, access_method: str) -> str:
    type_text = {
        "regulation": "法规正文、版本/修订状态、实施日期、公告及相关解释材料",
        "standard": "标准编号、名称、版本、状态、适用范围、引用关系及公开元数据",
        "certification": "认证制度、认可/认证机构信息、申请规则、技术文件要求及公开名单",
        "gma": "市场准入要求、产品合规义务、标签/能效/环保/贸易要求及办事指引",
    }.get(source_type, "公开法规认证资料")
    scope = str(crawl_scope or "").strip()
    method = str(access_method or "html").strip()
    prefix = f"可提取{type_text}"
    if scope:
        prefix += f"；重点范围：{scope}"
    return prefix + f"；接入形态：{method}。"


def _load_expansion_catalog() -> list[dict[str, Any]]:
    if not SOURCE_CATALOG_PATH.exists():
        return []
    payload = json.loads(SOURCE_CATALOG_PATH.read_text(encoding="utf-8"))
    if not isinstance(payload, list):
        raise ValueError("source expansion catalog must be a JSON list")
    return [dict(item) for item in payload if isinstance(item, dict)]


@dataclass(frozen=True)
class SeedSource:
    key: str
    region_code: str
    name: str
    source_type: str
    authority: str
    url: str
    access_method: str
    priority: str = "A"
    status: str = "research"
    shared_scope: bool = False
    languages: tuple[str, ...] = ()
    file_types: tuple[str, ...] = ("html",)
    refresh_policy: str = "daily"
    crawl_scope: str = ""
    notes: str = ""

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["languages"] = list(self.languages)
        data["file_types"] = list(self.file_types)
        return data


SEED_SOURCES: tuple[SeedSource, ...] = (
    # EU shared sources used by EU member markets.
    SeedSource("EU-EURLEX", "EU", "EUR-Lex", "regulation", "European Union", "https://eur-lex.europa.eu/", "mixed", shared_scope=True, languages=("multi",), file_types=("html", "pdf", "xml"), crawl_scope="EU legislation, directives, regulations, decisions and consolidated versions", notes="EU common legal source"),
    SeedSource("EU-ECHA", "EU", "ECHA", "gma", "European Chemicals Agency", "https://echa.europa.eu/", "mixed", shared_scope=True, languages=("en",), file_types=("html", "pdf", "table"), crawl_scope="REACH, CLP, SVHC and product chemical compliance information"),
    SeedSource("EU-EPREL", "EU", "EPREL", "gma", "European Commission", "https://eprel.ec.europa.eu/", "mixed", shared_scope=True, languages=("multi",), file_types=("html", "pdf", "structured"), crawl_scope="energy label and product registration information"),
    SeedSource("EU-ACCESS2MARKETS", "EU", "Access2Markets", "gma", "European Commission", "https://trade.ec.europa.eu/access-to-markets/en/home", "mixed", shared_scope=True, languages=("multi",), file_types=("html",), crawl_scope="market access requirements and trade procedures"),
    SeedSource("EU-HARMONISED-STANDARDS", "EU", "European Commission Harmonised Standards", "standard", "European Commission", "https://single-market-economy.ec.europa.eu/single-market/european-standards/harmonised-standards_en", "mixed", shared_scope=True, languages=("en",), file_types=("html", "pdf"), crawl_scope="harmonised standards references published for EU legislation", notes="standard references/metadata rather than licensed full text"),

    # 21 target markets: primary legislation portals.
    SeedSource("DE-LAW", "DE", "Gesetze im Internet", "regulation", "Federal Ministry of Justice / Federal Office of Justice", "https://www.gesetze-im-internet.de/", "mixed", languages=("de",), file_types=("html", "pdf", "xml"), crawl_scope="federal laws and ordinances"),
    SeedSource("FR-LAW", "FR", "Légifrance", "regulation", "French Government", "https://www.legifrance.gouv.fr/", "mixed", languages=("fr",), file_types=("html",), crawl_scope="codes, laws, decrees and official journal"),
    SeedSource("IT-LAW", "IT", "Normattiva", "regulation", "Italian Government", "https://www.normattiva.it/", "mixed", languages=("it",), file_types=("html",), crawl_scope="national legislation and versions"),
    SeedSource("ES-LAW", "ES", "BOE Legislación", "regulation", "Agencia Estatal Boletín Oficial del Estado", "https://www.boe.es/buscar/legislacion.php", "mixed", languages=("es",), file_types=("html", "pdf", "xml"), crawl_scope="state and selected EU/autonomous legislation"),
    SeedSource("NL-LAW", "NL", "Wetten.nl", "regulation", "Dutch Government / KOOP", "https://wetten.overheid.nl/", "xml", languages=("nl",), file_types=("html", "xml"), crawl_scope="national laws and regulations with historical versions"),
    SeedSource("BE-LAW", "BE", "Belgian Legislation Portal", "regulation", "Belgian Federal Public Service Justice", "https://www.ejustice.just.fgov.be/cgi_loi/welcome.pl", "html", languages=("nl", "fr", "de"), file_types=("html",), crawl_scope="federal legislation"),
    SeedSource("SE-LAW", "SE", "Sveriges riksdag - Dokument & lagar", "regulation", "Sveriges riksdag", "https://www.riksdagen.se/sv/dokument-och-lagar/", "mixed", languages=("sv",), file_types=("html", "open_data"), crawl_scope="laws, ordinances and parliamentary documents"),
    SeedSource("DK-LAW", "DK", "Retsinformation", "regulation", "Danish Government", "https://www.retsinformation.dk/", "html", languages=("da",), file_types=("html",), crawl_scope="laws and executive orders"),
    SeedSource("FI-LAW", "FI", "Finlex", "regulation", "Ministry of Justice / Legal Register Centre", "https://www.finlex.fi/en/legislation", "mixed", languages=("fi", "sv", "en"), file_types=("html", "pdf", "open_data"), crawl_scope="up-to-date legislation, statute book and amendments"),
    SeedSource("AT-LAW", "AT", "RIS Legal Information System", "regulation", "Federal Chancellery of Austria", "https://www.ris.bka.gv.at/", "mixed", languages=("de", "en"), file_types=("html", "open_data"), crawl_scope="federal/state legislation, gazettes and case law"),
    SeedSource("IE-LAW", "IE", "Irish Statute Book", "regulation", "Government of Ireland", "https://www.irishstatutebook.ie/", "html", languages=("en", "ga"), file_types=("html",), crawl_scope="Acts and statutory instruments"),
    SeedSource("NO-LAW", "NO", "Lovdata", "regulation", "Lovdata Foundation", "https://lovdata.no/", "mixed", languages=("no", "en"), file_types=("html",), crawl_scope="acts, central regulations and EEA agreement"),
    SeedSource("CH-LAW", "CH", "Fedlex", "regulation", "Swiss Confederation", "https://www.fedlex.admin.ch/", "mixed", languages=("de", "fr", "it", "rm", "en"), file_types=("html", "structured"), crawl_scope="federal law, official compilation and classified compilation"),
    SeedSource("GB-LAW", "GB", "legislation.gov.uk", "regulation", "The National Archives", "https://www.legislation.gov.uk/", "mixed", languages=("en",), file_types=("html", "xml", "pdf"), crawl_scope="UK primary and secondary legislation"),
    SeedSource("US-ECFR", "US", "eCFR", "regulation", "US Government Publishing Office / Office of the Federal Register", "https://www.ecfr.gov/", "mixed", languages=("en",), file_types=("html", "structured"), crawl_scope="current federal regulations"),
    SeedSource("US-FEDREG", "US", "Federal Register", "regulation", "Office of the Federal Register", "https://www.federalregister.gov/", "api", languages=("en",), file_types=("html", "json", "pdf"), crawl_scope="proposed and final federal rules, notices and presidential documents"),
    SeedSource("CA-LAW", "CA", "Justice Laws Website", "regulation", "Department of Justice Canada", "https://laws-lois.justice.gc.ca/", "mixed", languages=("en", "fr"), file_types=("html", "pdf"), crawl_scope="official consolidated federal Acts and regulations"),
    SeedSource("JP-LAW", "JP", "e-Gov 法令検索", "regulation", "Digital Agency / Government of Japan", "https://elaws.e-gov.go.jp/", "mixed", languages=("ja",), file_types=("html", "xml"), crawl_scope="Japanese laws, cabinet orders and ministerial ordinances"),
    SeedSource("KR-LAW", "KR", "国家法令信息中心", "regulation", "Ministry of Government Legislation", "https://www.law.go.kr/", "mixed", languages=("ko", "en"), file_types=("html",), crawl_scope="laws, presidential decrees, ministerial ordinances and administrative rules"),
    SeedSource("AU-LAW", "AU", "Federal Register of Legislation", "regulation", "Office of Parliamentary Counsel", "https://www.legislation.gov.au/", "mixed", languages=("en",), file_types=("html", "pdf", "structured"), crawl_scope="Commonwealth Acts and legislative instruments"),
    SeedSource("NZ-LAW", "NZ", "New Zealand Legislation", "regulation", "Parliamentary Counsel Office", "https://www.legislation.govt.nz/", "mixed", languages=("en",), file_types=("html", "pdf", "structured"), crawl_scope="Acts, Bills and secondary legislation"),
    SeedSource("SG-LAW", "SG", "Singapore Statutes Online", "regulation", "Attorney-General's Chambers", "https://sso.agc.gov.sg/", "mixed", languages=("en",), file_types=("html", "pdf", "rss"), crawl_scope="Acts and subsidiary legislation"),

    # Standards bodies / catalogues. Full text is subject to licence; metadata is the baseline.
    SeedSource("DE-STD", "DE", "DIN", "standard", "DIN Deutsches Institut für Normung", "https://www.din.de/", "metadata_only", languages=("de", "en"), file_types=("metadata",), refresh_policy="weekly", crawl_scope="DIN/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("FR-STD", "FR", "AFNOR", "standard", "AFNOR", "https://www.afnor.org/", "metadata_only", languages=("fr", "en"), file_types=("metadata",), refresh_policy="weekly", crawl_scope="NF/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("IT-STD", "IT", "UNI", "standard", "Ente Italiano di Normazione", "https://www.uni.com/", "metadata_only", languages=("it",), file_types=("metadata",), refresh_policy="weekly", crawl_scope="UNI/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("ES-STD", "ES", "UNE", "standard", "Asociación Española de Normalización", "https://www.une.org/", "metadata_only", languages=("es",), file_types=("metadata",), refresh_policy="weekly", crawl_scope="UNE/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("NL-STD", "NL", "NEN", "standard", "NEN", "https://www.nen.nl/", "metadata_only", languages=("nl", "en"), file_types=("metadata",), refresh_policy="weekly", crawl_scope="NEN/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("BE-STD", "BE", "NBN", "standard", "Bureau for Standardisation", "https://www.nbn.be/en", "metadata_only", languages=("nl", "fr", "en"), file_types=("metadata",), refresh_policy="weekly", crawl_scope="Belgian and European standard metadata", notes="full text subject to licence"),
    SeedSource("SE-STD", "SE", "SIS", "standard", "Swedish Institute for Standards", "https://www.sis.se/en/", "metadata_only", languages=("sv", "en"), file_types=("metadata",), refresh_policy="weekly", crawl_scope="SS/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("FI-STD", "FI", "SFS", "standard", "SFS Finnish Standards", "https://sfs.fi/en/", "metadata_only", languages=("fi", "en"), file_types=("metadata",), refresh_policy="weekly", crawl_scope="SFS/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("AT-STD", "AT", "Austrian Standards", "standard", "Austrian Standards International", "https://www.austrian-standards.at/en", "metadata_only", languages=("de", "en"), file_types=("metadata",), refresh_policy="weekly", crawl_scope="ÖNORM/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("IE-STD", "IE", "NSAI Standards", "standard", "National Standards Authority of Ireland", "https://www.nsai.ie/standards/", "metadata_only", languages=("en",), file_types=("metadata",), refresh_policy="weekly", crawl_scope="Irish/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("CH-STD", "CH", "SNV", "standard", "Swiss Association for Standardization", "https://www.snv.ch/", "metadata_only", languages=("de", "fr", "it"), file_types=("metadata",), refresh_policy="weekly", crawl_scope="SN/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("GB-STD", "GB", "BSI", "standard", "British Standards Institution", "https://www.bsigroup.com/en-GB/about-bsi/national-standards-body/", "metadata_only", languages=("en",), file_types=("metadata",), refresh_policy="weekly", crawl_scope="BS/EN/ISO standard metadata", notes="full text subject to licence"),
    SeedSource("US-STD", "US", "ANSI", "standard", "American National Standards Institute", "https://www.ansi.org/", "metadata_only", languages=("en",), file_types=("metadata",), refresh_policy="weekly", crawl_scope="American National Standards metadata and conformity assessment references", notes="full text depends on standards developer licence"),
    SeedSource("KR-STD-CERT", "KR", "KATS", "certification", "Korean Agency for Technology and Standards", "https://www.kats.go.kr/en/", "mixed", languages=("ko", "en"), file_types=("html",), refresh_policy="weekly", crawl_scope="Korean standards policy, KS certification and technical regulations"),
    SeedSource("AU-STD", "AU", "Standards Australia", "standard", "Standards Australia", "https://www.standards.org.au/", "metadata_only", languages=("en",), file_types=("metadata",), refresh_policy="weekly", crawl_scope="Australian Standards catalogue", notes="full text subject to licence"),
)


def _preclassify_harvestability(item: dict[str, Any]) -> tuple[str, str]:
    access_method = str(item.get("access_method") or "html")
    source_type = str(item.get("source_type") or "")
    notes = (str(item.get("notes") or "") + " " + str(item.get("extractable_summary") or "")).lower()
    if access_method == "metadata_only":
        return "metadata_only", "标准/目录类来源按公开元数据采集，标准全文仍按授权范围处理。"
    if source_type == "standard" and any(token in notes for token in ("licence", "license", "版权", "授权")):
        return "manual_authorized", "存在标准全文版权或许可边界，优先自动采集公开元数据，全文需授权后处理。"
    if access_method in {"html", "pdf", "xml", "api", "rss"}:
        return "direct", f"按{access_method}入口具备直接采集条件，仍需联网实测确认访问和内容结构。"
    if access_method == "mixed":
        return "adapter", "来源包含多种页面或资料形态，建议先联网核验，再配置站点级采集适配。"
    return "adapter", "当前接入方式需要进一步确认并配置采集适配。"


class SourceRegistryService:
    def __init__(self, store: KnowledgeStore):
        self.store = store
        self._ensure_schema()
        self._seed()
        self._ensure_preclassification()

    def _ensure_schema(self) -> None:
        with self.store.lock:
            self.store.conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS knowledge_sources (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    source_key TEXT NOT NULL UNIQUE,
                    region_code TEXT NOT NULL,
                    source_name TEXT NOT NULL,
                    source_type TEXT NOT NULL,
                    authority TEXT DEFAULT '',
                    base_url TEXT NOT NULL,
                    access_method TEXT DEFAULT 'html',
                    priority TEXT DEFAULT 'B',
                    status TEXT DEFAULT 'research',
                    shared_scope INTEGER DEFAULT 0,
                    languages_json TEXT DEFAULT '[]',
                    file_types_json TEXT DEFAULT '[]',
                    refresh_policy TEXT DEFAULT 'weekly',
                    harvestability TEXT DEFAULT '',
                    verification_status TEXT DEFAULT 'pending',
                    last_verified_at TEXT DEFAULT '',
                    last_http_status INTEGER DEFAULT 0,
                    last_content_type TEXT DEFAULT '',
                    last_final_url TEXT DEFAULT '',
                    robots_allowed TEXT DEFAULT 'unknown',
                    verification_note TEXT DEFAULT '',
                    source_summary TEXT DEFAULT '',
                    extractable_summary TEXT DEFAULT '',
                    crawl_scope TEXT DEFAULT '',
                    notes TEXT DEFAULT '',
                    owner_note TEXT DEFAULT '',
                    last_checked_at TEXT DEFAULT '',
                    last_success_at TEXT DEFAULT '',
                    created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT DEFAULT CURRENT_TIMESTAMP
                );
                CREATE INDEX IF NOT EXISTS idx_knowledge_sources_region ON knowledge_sources(region_code);
                CREATE INDEX IF NOT EXISTS idx_knowledge_sources_type ON knowledge_sources(source_type);
                """
            )
            columns = {
                row["name"]
                for row in self.store.conn.execute("PRAGMA table_info(knowledge_sources)").fetchall()
            }
            if "source_summary" not in columns:
                self.store.conn.execute("ALTER TABLE knowledge_sources ADD COLUMN source_summary TEXT DEFAULT ''")
            if "extractable_summary" not in columns:
                self.store.conn.execute("ALTER TABLE knowledge_sources ADD COLUMN extractable_summary TEXT DEFAULT ''")
            migrations = {
                "harvestability": "TEXT DEFAULT ''",
                "verification_status": "TEXT DEFAULT 'pending'",
                "last_verified_at": "TEXT DEFAULT ''",
                "last_http_status": "INTEGER DEFAULT 0",
                "last_content_type": "TEXT DEFAULT ''",
                "last_final_url": "TEXT DEFAULT ''",
                "robots_allowed": "TEXT DEFAULT 'unknown'",
                "verification_note": "TEXT DEFAULT ''",
            }
            for column, definition in migrations.items():
                if column not in columns:
                    self.store.conn.execute(f"ALTER TABLE knowledge_sources ADD COLUMN {column} {definition}")
            self.store.conn.commit()

    def _seed(self) -> None:
        with self.store.lock:
            for seed in SEED_SOURCES:
                item = seed.to_dict()
                source_summary = _default_source_summary(
                    name=item["name"],
                    authority=item["authority"],
                    source_type=item["source_type"],
                    crawl_scope=item["crawl_scope"],
                )
                extractable_summary = _default_extractable_summary(
                    source_type=item["source_type"],
                    crawl_scope=item["crawl_scope"],
                    access_method=item["access_method"],
                )
                self.store.conn.execute(
                    """
                    INSERT INTO knowledge_sources(
                        source_key,region_code,source_name,source_type,authority,base_url,access_method,
                        priority,status,shared_scope,languages_json,file_types_json,refresh_policy,
                        source_summary,extractable_summary,crawl_scope,notes
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(source_key) DO UPDATE SET
                        source_summary=CASE WHEN knowledge_sources.source_summary='' THEN excluded.source_summary ELSE knowledge_sources.source_summary END,
                        extractable_summary=CASE WHEN knowledge_sources.extractable_summary='' THEN excluded.extractable_summary ELSE knowledge_sources.extractable_summary END
                    """,
                    (
                        item["key"], canonical_region_code(item["region_code"]), item["name"], item["source_type"],
                        item["authority"], item["url"], item["access_method"], item["priority"], item["status"],
                        1 if item["shared_scope"] else 0, json.dumps(item["languages"], ensure_ascii=False),
                        json.dumps(item["file_types"], ensure_ascii=False), item["refresh_policy"],
                        source_summary, extractable_summary, item["crawl_scope"], item["notes"],
                    ),
                )

            for item in _load_expansion_catalog():
                source_key = str(item.get("source_key") or "").strip()
                region_code = canonical_region_code(str(item.get("region_code") or "").strip())
                source_name = str(item.get("source_name") or "").strip()
                source_type = str(item.get("source_type") or "").strip()
                authority = str(item.get("authority") or "").strip()
                base_url = str(item.get("base_url") or "").strip()
                access_method = str(item.get("access_method") or "html").strip()
                priority = str(item.get("priority") or "B").strip().upper()
                status = str(item.get("status") or "research").strip()
                shared_scope = 1 if bool(item.get("shared_scope")) else 0
                languages = item.get("languages") or []
                file_types = item.get("file_types") or ["html"]
                refresh_policy = str(item.get("refresh_policy") or "weekly").strip()
                crawl_scope = str(item.get("crawl_scope") or "").strip()
                notes = str(item.get("notes") or "").strip()
                source_summary = str(item.get("source_summary") or "").strip() or _default_source_summary(
                    name=source_name,
                    authority=authority,
                    source_type=source_type,
                    crawl_scope=crawl_scope,
                )
                extractable_summary = str(item.get("extractable_summary") or "").strip() or _default_extractable_summary(
                    source_type=source_type,
                    crawl_scope=crawl_scope,
                    access_method=access_method,
                )
                self.store.conn.execute(
                    """
                    INSERT INTO knowledge_sources(
                        source_key,region_code,source_name,source_type,authority,base_url,access_method,
                        priority,status,shared_scope,languages_json,file_types_json,refresh_policy,
                        source_summary,extractable_summary,crawl_scope,notes
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(source_key) DO UPDATE SET
                        source_summary=CASE WHEN knowledge_sources.source_summary='' THEN excluded.source_summary ELSE knowledge_sources.source_summary END,
                        extractable_summary=CASE WHEN knowledge_sources.extractable_summary='' THEN excluded.extractable_summary ELSE knowledge_sources.extractable_summary END
                    """,
                    (
                        source_key, region_code, source_name, source_type, authority, base_url, access_method,
                        priority, status, shared_scope, json.dumps(languages, ensure_ascii=False),
                        json.dumps(file_types, ensure_ascii=False), refresh_policy, source_summary,
                        extractable_summary, crawl_scope, notes,
                    ),
                )
            self.store.conn.commit()

    def _ensure_preclassification(self) -> None:
        with self.store.lock:
            rows = self.store.conn.execute(
                "SELECT * FROM knowledge_sources WHERE harvestability='' OR verification_status='pending'"
            ).fetchall()
            for row in rows:
                item = dict(row)
                harvestability, note = _preclassify_harvestability(item)
                self.store.conn.execute(
                    """UPDATE knowledge_sources
                       SET harvestability=?,
                           verification_status=CASE WHEN verification_status='pending' THEN 'preclassified' ELSE verification_status END,
                           verification_note=CASE WHEN verification_note='' THEN ? ELSE verification_note END
                       WHERE id=?""",
                    (harvestability, note, int(item["id"])),
                )
            self.store.conn.commit()

    def record_verification(self, source_id: int, payload: dict[str, Any]) -> dict[str, Any]:
        harvestability = str(payload.get("harvestability") or "").strip()
        verification_status = str(payload.get("verification_status") or "").strip()
        if harvestability not in HARVESTABILITY_TYPES:
            raise ValueError(f"unsupported harvestability: {harvestability}")
        if verification_status not in VERIFICATION_STATUSES:
            raise ValueError(f"unsupported verification status: {verification_status}")
        params = (
            harvestability,
            verification_status,
            str(payload.get("last_verified_at") or ""),
            int(payload.get("last_http_status") or 0),
            str(payload.get("last_content_type") or ""),
            str(payload.get("last_final_url") or ""),
            str(payload.get("robots_allowed") or "unknown"),
            str(payload.get("verification_note") or ""),
            int(source_id),
        )
        with self.store.lock:
            cur = self.store.conn.execute(
                """UPDATE knowledge_sources
                   SET harvestability=?,verification_status=?,last_verified_at=?,
                       last_http_status=?,last_content_type=?,last_final_url=?,
                       robots_allowed=?,verification_note=?,updated_at=CURRENT_TIMESTAMP
                   WHERE id=?""",
                params,
            )
            if not cur.rowcount:
                raise ValueError("source not found")
            self.store.conn.commit()
        return self.detail(source_id)

    def verification_summary(self) -> dict[str, Any]:
        with self.store.lock:
            rows = self.store.conn.execute(
                """SELECT harvestability,verification_status,COUNT(*) count
                   FROM knowledge_sources
                   WHERE status!='archived'
                   GROUP BY harvestability,verification_status"""
            ).fetchall()
        by_harvestability: Counter[str] = Counter()
        by_verification: Counter[str] = Counter()
        for row in rows:
            count = int(row["count"] or 0)
            by_harvestability[str(row["harvestability"] or "unknown")] += count
            by_verification[str(row["verification_status"] or "pending")] += count
        return {
            "total": sum(by_harvestability.values()),
            "by_harvestability": dict(by_harvestability),
            "by_verification_status": dict(by_verification),
        }

    def _row(self, row) -> dict[str, Any]:
        item = dict(row)
        item["shared_scope"] = bool(item.get("shared_scope"))
        item["languages"] = json.loads(item.pop("languages_json") or "[]")
        item["file_types"] = json.loads(item.pop("file_types_json") or "[]")
        return item

    def list_sources(self, *, region_code: str = "", source_type: str = "", status: str = "", harvestability: str = "", verification_status: str = "", q: str = "", limit: int = 1000) -> dict[str, Any]:
        clauses: list[str] = []
        params: list[Any] = []
        if region_code:
            clauses.append("region_code=?")
            params.append(canonical_region_code(region_code))
        if source_type:
            clauses.append("source_type=?")
            params.append(source_type)
        if status:
            clauses.append("status=?")
            params.append(status)
        if harvestability:
            if harvestability not in HARVESTABILITY_TYPES:
                raise ValueError(f"unsupported harvestability: {harvestability}")
            clauses.append("harvestability=?")
            params.append(harvestability)
        if verification_status:
            if verification_status not in VERIFICATION_STATUSES:
                raise ValueError(f"unsupported verification status: {verification_status}")
            clauses.append("verification_status=?")
            params.append(verification_status)
        if q:
            clauses.append("(lower(source_key) LIKE ? OR lower(source_name) LIKE ? OR lower(authority) LIKE ? OR lower(base_url) LIKE ? OR lower(source_summary) LIKE ? OR lower(extractable_summary) LIKE ? OR lower(crawl_scope) LIKE ?)")
            term = f"%{q.lower()}%"
            params.extend([term, term, term, term, term, term, term])
        where = "WHERE " + " AND ".join(clauses) if clauses else ""
        params.append(min(max(int(limit), 1), 3000))
        with self.store.lock:
            rows = self.store.conn.execute(
                f"SELECT * FROM knowledge_sources {where} ORDER BY region_code,priority,source_type,source_name LIMIT ?",
                params,
            ).fetchall()
        items = [self._row(row) for row in rows]
        return {"items": items, "summary": self.summary(items=items)}

    def summary(self, *, items: list[dict[str, Any]] | None = None) -> dict[str, Any]:
        if items is None:
            with self.store.lock:
                rows = self.store.conn.execute("SELECT * FROM knowledge_sources").fetchall()
            items = [self._row(row) for row in rows]
        by_type = Counter(item.get("source_type") for item in items)
        by_status = Counter(item.get("status") for item in items)
        active_items = [item for item in items if item.get("status") != "archived"]
        active_by_type = Counter(item.get("source_type") for item in active_items)
        regions = {item.get("region_code") for item in active_items if item.get("region_code") not in {"", "EU"}}
        target_codes = {market.code for market in TARGET_MARKETS}
        return {
            "target_sources": TARGET_SOURCE_COUNT,
            "registered_sources": len(active_items),
            "archived_sources": len(items) - len(active_items),
            "remaining_to_target": max(0, TARGET_SOURCE_COUNT - len(active_items)),
            "target_markets": len(TARGET_MARKETS),
            "target_markets_with_sources": len(regions & target_codes),
            "by_type": dict(active_by_type),
            "by_status": dict(by_status),
            "by_harvestability": dict(Counter(item.get("harvestability") or "unknown" for item in active_items)),
            "by_verification_status": dict(Counter(item.get("verification_status") or "pending" for item in active_items)),
            "shared_sources": sum(1 for item in items if item.get("shared_scope")),
        }

    def by_key(self, source_key: str) -> dict[str, Any]:
        key = str(source_key or "").strip()
        if not key:
            raise ValueError("source_key is required")
        with self.store.lock:
            row = self.store.conn.execute(
                "SELECT * FROM knowledge_sources WHERE source_key=?", (key,)
            ).fetchone()
        if not row:
            raise ValueError(f"source not found: {key}")
        return self._row(row)

    @staticmethod
    def _list_value(value: Any) -> list[str]:
        if isinstance(value, (list, tuple, set)):
            return [str(item).strip() for item in value if str(item).strip()]
        text = str(value or "").strip()
        if not text:
            return []
        for delimiter in ("，", "；", ";", "|", "、"):
            text = text.replace(delimiter, ",")
        return [item.strip() for item in text.split(",") if item.strip()]

    @staticmethod
    def _bool_value(value: Any) -> bool:
        if isinstance(value, bool):
            return value
        return str(value or "").strip().lower() in {"1", "true", "yes", "y", "是", "共享"}

    def _normalize_payload(self, payload: dict[str, Any], *, creating: bool = False) -> dict[str, Any]:
        result: dict[str, Any] = {}
        target_codes = {market.code for market in TARGET_MARKETS} | {"EU"}

        if creating or "source_key" in payload:
            source_key = str(payload.get("source_key") or "").strip()
            if not source_key:
                raise ValueError("source_key is required")
            result["source_key"] = source_key

        if creating or "region_code" in payload:
            region_code = canonical_region_code(str(payload.get("region_code") or "").strip())
            if not region_code or region_code not in target_codes:
                raise ValueError(f"unsupported region_code: {region_code or payload.get('region_code')}")
            result["region_code"] = region_code

        if creating or "source_name" in payload:
            source_name = str(payload.get("source_name") or "").strip()
            if not source_name:
                raise ValueError("source_name is required")
            result["source_name"] = source_name

        if creating or "source_type" in payload:
            raw = str(payload.get("source_type") or "").strip()
            source_type = SOURCE_TYPE_ALIASES.get(raw, raw)
            if source_type not in SOURCE_TYPES:
                raise ValueError(f"unsupported source_type: {raw}")
            result["source_type"] = source_type

        if creating or "base_url" in payload:
            base_url = str(payload.get("base_url") or "").strip()
            if not base_url:
                raise ValueError("base_url is required")
            if not (base_url.startswith("http://") or base_url.startswith("https://")):
                raise ValueError("base_url must start with http:// or https://")
            result["base_url"] = base_url

        text_fields = ("authority", "source_summary", "extractable_summary", "crawl_scope", "notes", "owner_note")
        for key in text_fields:
            if key in payload:
                result[key] = str(payload.get(key) or "").strip()

        if "access_method" in payload or creating:
            raw = str(payload.get("access_method") or "html").strip()
            access_method = ACCESS_METHOD_ALIASES.get(raw, raw)
            if access_method not in ACCESS_METHODS:
                raise ValueError(f"unsupported access method: {raw}")
            result["access_method"] = access_method

        if "priority" in payload or creating:
            priority = str(payload.get("priority") or "B").strip().upper()
            if priority not in SOURCE_PRIORITIES:
                raise ValueError(f"unsupported priority: {priority}")
            result["priority"] = priority

        if "status" in payload or creating:
            raw = str(payload.get("status") or "research").strip()
            status = SOURCE_STATUS_ALIASES.get(raw, raw)
            if status not in SOURCE_STATUSES:
                raise ValueError(f"unsupported source status: {raw}")
            result["status"] = status

        if "refresh_policy" in payload or creating:
            raw = str(payload.get("refresh_policy") or "weekly").strip()
            refresh_policy = REFRESH_POLICY_ALIASES.get(raw, raw)
            if refresh_policy not in REFRESH_POLICIES:
                raise ValueError(f"unsupported refresh policy: {raw}")
            result["refresh_policy"] = refresh_policy

        if "shared_scope" in payload:
            result["shared_scope"] = 1 if self._bool_value(payload.get("shared_scope")) else 0
        elif creating:
            result["shared_scope"] = 0

        if "languages" in payload:
            result["languages_json"] = json.dumps(self._list_value(payload.get("languages")), ensure_ascii=False)
        elif creating:
            result["languages_json"] = "[]"

        if "file_types" in payload:
            result["file_types_json"] = json.dumps(self._list_value(payload.get("file_types")), ensure_ascii=False)
        elif creating:
            result["file_types_json"] = "[]"

        return result

    def detail(self, source_id: int) -> dict[str, Any]:
        with self.store.lock:
            row = self.store.conn.execute("SELECT * FROM knowledge_sources WHERE id=?", (int(source_id),)).fetchone()
        if not row:
            raise ValueError("source not found")
        return self._row(row)

    def create(self, payload: dict[str, Any]) -> dict[str, Any]:
        item = self._normalize_payload(payload, creating=True)
        columns = [
            "source_key", "region_code", "source_name", "source_type", "authority", "base_url",
            "access_method", "priority", "status", "shared_scope", "languages_json", "file_types_json",
            "refresh_policy", "source_summary", "extractable_summary", "crawl_scope", "notes", "owner_note",
        ]
        values = [item.get(key, "") for key in columns]
        with self.store.lock:
            existing = self.store.conn.execute(
                "SELECT id FROM knowledge_sources WHERE source_key=?",
                (item["source_key"],),
            ).fetchone()
            if existing:
                raise ValueError(f"source_key already exists: {item['source_key']}")
            cur = self.store.conn.execute(
                f"INSERT INTO knowledge_sources({','.join(columns)}) VALUES({','.join('?' for _ in columns)})",
                values,
            )
            self.store.conn.commit()
            source_id = int(cur.lastrowid)
        return self.detail(source_id)

    def update(self, source_id: int, payload: dict[str, Any]) -> dict[str, Any]:
        item = self._normalize_payload(payload, creating=False)
        item.pop("source_key", None)
        if not item:
            return self.detail(source_id)
        updates = [f"{key}=?" for key in item]
        params = list(item.values()) + [int(source_id)]
        with self.store.lock:
            cur = self.store.conn.execute(
                f"UPDATE knowledge_sources SET {','.join(updates)},updated_at=CURRENT_TIMESTAMP WHERE id=?",
                params,
            )
            if not cur.rowcount:
                raise ValueError("source not found")
            self.store.conn.commit()
        return self.detail(source_id)

    def archive(self, source_id: int, *, note: str = "") -> dict[str, Any]:
        payload: dict[str, Any] = {"status": "archived"}
        if note:
            current = self.detail(source_id)
            owner_note = str(current.get("owner_note") or "").strip()
            payload["owner_note"] = (owner_note + "\n" + note).strip()
        return self.update(source_id, payload)

    def batch_update(self, source_ids: list[int], payload: dict[str, Any]) -> dict[str, Any]:
        ids = sorted({int(source_id) for source_id in source_ids if int(source_id) > 0})
        if not ids:
            raise ValueError("source_ids is required")
        allowed = {key: payload[key] for key in ("status", "priority") if key in payload}
        if not allowed:
            raise ValueError("batch update requires status or priority")
        normalized = self._normalize_payload(allowed, creating=False)
        updates = [f"{key}=?" for key in normalized]
        placeholders = ",".join("?" for _ in ids)
        params = list(normalized.values()) + ids
        with self.store.lock:
            cur = self.store.conn.execute(
                f"UPDATE knowledge_sources SET {','.join(updates)},updated_at=CURRENT_TIMESTAMP WHERE id IN ({placeholders})",
                params,
            )
            self.store.conn.commit()
        return {"updated": int(cur.rowcount or 0), "source_ids": ids}

    def export_xlsx(self, *, template_only: bool = False) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "来源台账"
        sheet.append([label for label, _ in SOURCE_EXPORT_COLUMNS])
        if not template_only:
            items = self.list_sources(limit=3000)["items"]
            for item in items:
                row = []
                for _, key in SOURCE_EXPORT_COLUMNS:
                    value: Any = item.get(key, "")
                    if key in {"languages", "file_types"}:
                        value = "、".join(value or [])
                    elif key == "shared_scope":
                        value = "是" if value else "否"
                    row.append(value)
                sheet.append(row)
        widths = [18, 12, 28, 16, 28, 44, 14, 10, 12, 10, 18, 18, 12, 16, 14, 20, 12, 24, 48, 48, 52, 42, 36, 36]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
        sheet.freeze_panes = "A2"
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def import_xlsx(self, content: bytes) -> dict[str, Any]:
        if not content:
            raise ValueError("empty workbook")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            raise ValueError("workbook is empty")
        mapped = [SOURCE_HEADER_ALIASES.get(str(value or "").strip(), "") for value in headers]
        if "source_key" not in mapped or "source_name" not in mapped or "base_url" not in mapped:
            raise ValueError("Excel must include 来源编码、来源名称、网址")

        created = 0
        updated = 0
        errors: list[dict[str, Any]] = []
        for row_number, values in enumerate(rows, start=2):
            payload = {
                key: values[index]
                for index, key in enumerate(mapped)
                if key and index < len(values) and values[index] not in {None, ""}
            }
            if not payload:
                continue
            source_key = str(payload.get("source_key") or "").strip()
            try:
                if not source_key:
                    raise ValueError("来源编码不能为空")
                try:
                    existing = self.by_key(source_key)
                except ValueError:
                    existing = None
                if existing:
                    self.update(int(existing["id"]), payload)
                    updated += 1
                else:
                    self.create(payload)
                    created += 1
            except Exception as exc:
                errors.append({"row": row_number, "source_key": source_key, "error": str(exc)})
        return {
            "created": created,
            "updated": updated,
            "failed": len(errors),
            "errors": errors[:100],
        }
